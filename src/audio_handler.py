# region Imports

import os
import glob
import math
import sys
from pydub import AudioSegment
import openpyxl
from typing import List
import src.print_color as pc
from src.logger_setup import gui_logger as logger
from src.audio_config import SPLIT_CONFIG

# endregion Imports


# region Audio Processing

class AudioSplitter:
    def __init__(self):
        self.temp_dir = None

    def __enter__(self):
        self.temp_dir = self.create_temp_dir()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.clear_temp_dir()

    def split_audio(
        self, file: str, max_size_mb: int = SPLIT_CONFIG["max_file_size"], overlap_ms: int = SPLIT_CONFIG["overlap_ms"]
    ) -> List[tuple]:
        logger.info(f"Splitting audio file: {file}")
        split_files = []
        segment_count = self.get_segment_count(file, max_size_mb)
        if segment_count > 1:
            song = AudioSegment.from_file(file)
            pc.print_info(f"\nsplitting {file} into {segment_count} parts")
            segment_length_ms = len(song) / segment_count
            overlap = overlap_ms if segment_length_ms >= overlap_ms else 0

            for i in range(segment_count):
                seg_start_ms = (
                    math.floor(i * segment_length_ms - overlap / 2) if i > 0 else 0
                )
                seg_end_ms = (
                    math.ceil((i + 1) * segment_length_ms + overlap / 2)
                    if i < segment_count - 1
                    else None
                )

                segment = song[seg_start_ms:seg_end_ms]
                segment_filename = os.path.join(
                    self.temp_dir,
                    f"{os.path.splitext(os.path.basename(file))[0]}_segment{i}{os.path.splitext(file)[1]}",
                )
                segment.export(segment_filename)
                pc.print_info(f"\n  > saved segment {segment_filename}")
                split_files.append((segment_filename, seg_start_ms))
            pc.print_info(f"\nfinished splitting {file}")
        else:
            split_files.append((file, 0))
        return split_files

    def get_segment_count(self, file: str, max_size_mb: int = SPLIT_CONFIG["max_file_size"]) -> int:
        logger.debug(f"Calculating segments for file: {file}")
        
        max_size_bytes = max_size_mb * 1024 * 1024
        file_size = os.path.getsize(file)
        segments = math.ceil(file_size / max_size_bytes) if file_size > max_size_bytes else 1

        logger.debug(f"File size: {file_size / (1024 * 1024):.2f}MB, Segments needed: {segments}")
        return segments
    
    def get_temp_dir(self) -> str:
        logger.debug("Getting temp directory path")
        if getattr(sys, "frozen", False):
            return os.path.join(os.path.dirname(sys.executable), ".temp")
        elif __file__:
            return os.path.join(os.path.dirname(__file__), ".temp")
        return os.path.join(os.getcwd(), ".temp")

    def create_temp_dir(self) -> str:
        logger.info("Creating temporary directory")
        temp_dir = self.get_temp_dir()
        pc.print_info(f"\ntemp dir: {temp_dir}")
        if not os.path.exists(temp_dir):
            os.makedirs(temp_dir)
        return temp_dir

    def clear_temp_dir(self):
        logger.info("Clearing temporary directory")
        if self.temp_dir and os.path.exists(self.temp_dir):
            files = glob.glob(os.path.join(self.temp_dir, "*"))
            for f in files:
                os.remove(f)
            logger.info(f"Cleared {len(files)} files from temporary directory")
        else:
            logger.warning("Temporary directory does not exist or is not set")

# endregion Audio Processing


# region Transcription Handler

class TranscriptionHandler:
    @staticmethod
    def save_result(result: dict, output_formats: List[str], file_path: str):
        logger.info(f"Saving transcription results for {file_path}")
        output_dir = os.path.dirname(os.path.realpath(file_path))

        for ext in output_formats:
            output_filename = f"{os.path.splitext(file_path)[0]}.{ext}"
            logger.debug(f"Saving {ext} format to: {output_filename}")

            if ext == "xlsx":
                wb = openpyxl.Workbook()
                sheet = wb.active
                sheet.cell(row=1, column=1).value = "start"
                sheet.cell(row=1, column=2).value = "end"
                sheet.cell(row=1, column=3).value = "text"

                for i, segment in enumerate(result["segments"], 2):
                    sheet.cell(row=i, column=1).value = segment["start"]
                    sheet.cell(row=i, column=2).value = segment["end"]
                    sheet.cell(row=i, column=3).value = segment["text"]

                wb.save(os.path.join(output_dir, output_filename))
            else:
                # Handle other formats (txt, srt, vtt, etc.)
                with open(
                    os.path.join(output_dir, output_filename), "w", encoding="utf-8"
                ) as f:
                    if ext == "txt":
                        f.write(result["text"])
                    elif ext in ["srt", "vtt"]:
                        for i, segment in enumerate(result["segments"], 1):
                            f.write(f"{i}\n")
                            start = TranscriptionHandler.format_timestamp(
                                segment["start"]
                            )
                            end = TranscriptionHandler.format_timestamp(segment["end"])
                            f.write(f"{start} --> {end}\n")
                            f.write(f"{segment['text']}\n\n")

    @staticmethod
    def format_timestamp(seconds: float) -> str:
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        msecs = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{msecs:03d}"

# endregion Transcription Handler
